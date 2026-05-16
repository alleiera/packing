import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from settings_manager import SettingsManager

class ExcelExporter:
    @staticmethod
    def export(file_path, header_info, items, language='tr'):
        # Prepare data
        if language == 'tr':
            cols = ["Ürün Kodu", "Ürün Adı", "Ölçü", "Toplam Metre", "Koli Adeti", "Koli Ağırlığı", "Net Ağırlık", "Bürüt Ağırlık"]
        else:
            cols = ["Product Code", "Product Name", "Size", "Total Meter", "Box Count", "Box Weight", "Net Weight", "Gross Weight"]

        df = pd.DataFrame(items, columns=cols)

        # Create Excel with custom formatting to match template
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, startrow=10)

            workbook = writer.book
            worksheet = writer.sheets['Sheet1']

            # Add Company Info / Logo Placeholder (Real logo insertion needs PIL and openpyxl Image)
            settings = SettingsManager.get_all_settings()
            worksheet['A1'] = "PACKING LIST"
            worksheet['A1'].font = Font(bold=True, size=14)

            worksheet['A2'] = f"Document No: {header_info.get('doc_no', '')}"
            worksheet['A3'] = f"Date: {header_info.get('date', '')}"

            worksheet['A5'] = "CONSIGNEE"
            worksheet['A5'].font = Font(bold=True)
            worksheet['A6'] = header_info.get('consignee', '')

            worksheet['E1'] = settings.get('company_name', '')
            worksheet['E1'].font = Font(bold=True)
            worksheet['E2'] = settings.get('company_address', '')
            worksheet['E3'] = settings.get('company_tel', '')

            # Totals
            total_row = 10 + len(items) + 1
            worksheet.cell(row=total_row + 1, column=5, value="TOTAL BOX:")
            worksheet.cell(row=total_row + 1, column=6, value=header_info.get('total_boxes', 0))
            worksheet.cell(row=total_row + 2, column=5, value="TOTAL PALLET:")
            worksheet.cell(row=total_row + 2, column=6, value=header_info.get('total_pallets', 0))

        return True
